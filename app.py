import html
import io
import json
import os
import re
import urllib.parse
from datetime import datetime
import dateutil.parser
import feedparser
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Présidentielle 2027 — La course, en un coup d'œil",
    page_icon="⚖️",
    layout="wide",
)

DATE_MIN = datetime(2026, 1, 1, 0, 0, 0)
DB_LOIS = "veille_data.json"
DB_CANDIDATS = "candidats_presse.json"

# Mots-clés ciblés sur les déclarations de candidatures dans la presse française
QUERIES_CANDIDATS = [
    'présidentielle "se déclare candidat"',
    'présidentielle "annonce sa candidature"',
    'présidentielle "candidat officiel"',
    'présidentielle "candidature déclarée"',
    'présidentielle "investi par"',
    'présidentielle "sera candidat"',
]

# Mots-clés stricts : propositions de loi françaises & secteur agroalimentaire
QUERIES_LOIS = [
    'France "proposition de loi" présidentielle',
    '"proposition de loi" Assemblée nationale présidentielle',
    '"proposition de loi" Sénat présidentielle',
    'France "projet de loi" présidentielle',
    '"proposition de loi" agroalimentaire',
    '"proposition de loi" "taxe soda" OR "taxe sucre"',
    '"proposition de loi" agriculture alimentation',
    'présidentielle EGAlim négociations commerciales',
    '"proposition de loi" emballage plastique consigne',
    '"proposition de loi" grande distribution marges',
    'France "mesure fiscale" industrie agroalimentaire',
    'présidentielle fiscalité entreprises agroalimentaires',
]

PARTIS_MAPPING = [
    ("Horizons", "Horizons"),
    ("Renaissance", "Renaissance"),
    ("Rassemblement National", "Rassemblement National"),
    ("RN", "Rassemblement National"),
    ("La France Insoumise", "La France Insoumise"),
    ("LFI", "La France Insoumise"),
    ("Parti Socialiste", "Parti Socialiste"),
    ("PS", "Parti Socialiste"),
    ("Les Républicains", "Les Républicains"),
    ("LR", "Les Républicains"),
    ("Droite Républicaine", "La Droite Républicaine"),
    ("Écologistes", "Les Écologistes"),
    ("EELV", "Les Écologistes"),
    ("Parti Communiste", "Parti Communiste Français"),
    ("PCF", "Parti Communiste Français"),
    ("Reconquête", "Reconquête"),
    ("Debout", "Debout !"),
    ("MoDem", "MoDem"),
    ("UDI", "UDI"),
]

# Liste de personnalités politiques surveillées pour fiabiliser la détection
POLITIQUES_CONNUES = [
    ("Édouard Philippe", "Horizons"),
    ("Edouard Philippe", "Horizons"),
    ("Marine Le Pen", "Rassemblement National"),
    ("Jordan Bardella", "Rassemblement National"),
    ("Jean-Luc Mélenchon", "La France Insoumise"),
    ("François Ruffin", "Debout !"),
    ("Laurent Wauquiez", "La Droite Républicaine"),
    ("Gabriel Attal", "Renaissance"),
    ("Gérald Darmanin", "Renaissance"),
    ("Fabien Roussel", "Parti Communiste Français"),
    ("Bernard Cazeneuve", "La Convention"),
    ("David Lisnard", "Nouvelle Énergie"),
    ("Yannick Jadot", "Les Écologistes"),
    ("Marine Tondelier", "Les Écologistes"),
]

st.markdown(
    """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Lora:ital,wght@0,400;0,600;0,700;1,400;1,600&family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');

    .stApp {
        background-color: #F4F1EA;
        color: #1C1917;
        font-family: 'Plus Jakarta Sans', sans-serif;
    }

    .top-meta {
        display: flex;
        justify-content: space-between;
        font-size: 0.85rem;
        color: #57534E;
        margin-bottom: 18px;
    }
    .header-title {
        font-family: 'Lora', serif;
        font-size: 3.2rem;
        font-weight: 700;
        line-height: 1.15;
        letter-spacing: -0.5px;
        color: #1C1917;
        margin-bottom: 8px;
    }
    .header-title em {
        font-style: italic;
        font-weight: 600;
        color: #8B261E;
    }
    .header-lead {
        font-size: 1.05rem;
        color: #57534E;
        max-width: 780px;
        margin-bottom: 20px;
        line-height: 1.5;
    }
    .milestones {
        display: flex;
        gap: 24px;
        flex-wrap: wrap;
        font-size: 0.9rem;
        color: #57534E;
        padding-bottom: 16px;
    }
    .milestones strong {
        color: #1C1917;
        font-weight: 700;
    }
    .main-separator {
        border-top: 2px solid #1C1917;
        margin-bottom: 24px;
    }

    .candidate-card {
        background: rgba(255, 255, 255, 0.65);
        border: 1px solid #D6D3CD;
        border-left: 4px solid #8B261E;
        border-radius: 4px;
        padding: 18px 20px;
        margin-bottom: 14px;
    }
    .candidate-name {
        font-family: 'Lora', serif;
        font-weight: 700;
        font-size: 1.2rem;
        color: #1C1917;
    }
    .candidate-party {
        display: inline-block;
        background-color: #1C1917;
        color: #F4F1EA;
        font-size: 0.75rem;
        font-weight: 600;
        padding: 3px 8px;
        border-radius: 3px;
        margin: 6px 0;
        text-transform: uppercase;
        letter-spacing: 0.5px;
    }
    .candidate-quote {
        font-size: 0.88rem;
        color: #44403C;
        margin-top: 4px;
        line-height: 1.45;
    }
    .candidate-link {
        display: inline-block;
        margin-top: 6px;
        font-size: 0.8rem;
        color: #8B261E;
        text-decoration: none;
        font-weight: 600;
    }

    .stTabs [data-baseweb="tab-list"] {
        border-bottom: 1px solid #1C1917;
        gap: 8px;
    }
    .stTabs [data-baseweb="tab"] {
        font-family: 'Lora', serif;
        font-size: 1.1rem;
        font-weight: 600;
        color: #57534E;
        background-color: transparent;
        padding: 10px 20px;
        border: 1px solid transparent;
        border-bottom: none;
        border-radius: 4px 4px 0 0;
    }
    .stTabs [aria-selected="true"] {
        color: #1C1917 !important;
        background-color: #F4F1EA !important;
        border-color: #1C1917 !important;
        border-bottom: 1px solid #F4F1EA !important;
    }
</style>
""",
    unsafe_allow_html=True,
)


def parser_date(date_str):
    try:
        dt = dateutil.parser.parse(date_str)
        if dt.tzinfo is not None:
            dt = dt.replace(tzinfo=None)
        return dt
    except Exception:
        return None


def detecter_parti(texte):
    for declinaison, label in PARTIS_MAPPING:
        if re.search(r"\b" + re.escape(declinaison) + r"\b", texte, re.IGNORECASE):
            return label
    return "Mouvement politique en cours de précision"


def detecter_candidat(texte):
    for nom, parti_defaut in POLITIQUES_CONNUES:
        if re.search(r"\b" + re.escape(nom) + r"\b", texte, re.IGNORECASE):
            parti = detecter_parti(texte)
            return nom, parti if parti != "Mouvement politique en cours de précision" else parti_defaut
    return None, None


def charger_fichier(chemin):
    if os.path.exists(chemin):
        with open(chemin, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


def sauvegarder_fichier(chemin, data):
    with open(chemin, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# Synchronisation automatique dès l'ouverture de la page
@st.cache_data(ttl=1800, show_spinner=False)
def sync_automatique():
    # 1. Propositions de loi (Presse et institutions françaises uniquement)
    existants_lois = charger_fichier(DB_LOIS)
    vus_lois = {a["Lien"] for a in existants_lois}
    nouvelles_lois = []

    for q in QUERIES_LOIS:
        # Flux restreint à la France
        url = f"https://news.google.com/rss/search?q={urllib.parse.quote(q)}&hl=fr&gl=FR&ceid=FR:fr"
        flux = feedparser.parse(url)

        for entry in flux.entries:
            lien = entry.link
            if lien in vus_lois:
                continue

            dt = parser_date(entry.get("published", ""))
            if dt and dt < DATE_MIN:
                continue

            vus_lois.add(lien)
            source = entry.source.get("title", "Presse française") if hasattr(entry, "source") else "Presse"
            titre = html.unescape(entry.title)
            if " - " in titre:
                titre = titre.rsplit(" - ", 1)[0]

            nouvelles_lois.append({
                "Date": dt.strftime("%d %b %Y %H:%M") if dt else "2026",
                "Source": source,
                "Titre": titre,
                "Lien": lien,
                "_dt": dt.isoformat() if dt else "2026-01-01T00:00:00",
            })

    total_lois = nouvelles_lois + existants_lois
    # Plus récent en haut
    total_lois.sort(key=lambda x: x.get("_dt", ""), reverse=True)
    sauvegarder_fichier(DB_LOIS, total_lois)

    # 2. Candidats déclarés dans la presse
    existants_cand = charger_fichier(DB_CANDIDATS)
    candidats_map = {c["Nom"]: c for c in existants_cand}

    for q in QUERIES_CANDIDATS:
        url = f"https://news.google.com/rss/search?q={urllib.parse.quote(q)}&hl=fr&gl=FR&ceid=FR:fr"
        flux = feedparser.parse(url)

        for entry in flux.entries:
            dt = parser_date(entry.get("published", ""))
            if dt and dt < DATE_MIN:
                continue

            titre = html.unescape(entry.title)
            source = entry.source.get("title", "Presse") if hasattr(entry, "source") else "Presse"
            if " - " in titre:
                titre = titre.rsplit(" - ", 1)[0]

            nom, parti = detecter_candidat(titre)
            if nom:
                # Met à jour la dernière communication dans la presse pour ce candidat
                dt_iso = dt.isoformat() if dt else "2026-01-01T00:00:00"
                if nom not in candidats_map or dt_iso > candidats_map[nom].get("_dt", ""):
                    candidats_map[nom] = {
                        "Nom": nom,
                        "Parti": parti,
                        "DerniereAnnonce": titre,
                        "Source": source,
                        "Date": dt.strftime("%d %b %Y") if dt else "2026",
                        "Lien": entry.link,
                        "_dt": dt_iso,
                    }

    total_cand = list(candidats_map.values())
    total_cand.sort(key=lambda x: x.get("_dt", ""), reverse=True)
    sauvegarder_fichier(DB_CANDIDATS, total_cand)


# Déclenchement automatique
sync_automatique()

data_lois = charger_fichier(DB_LOIS)
data_cand = charger_fichier(DB_CANDIDATS)


def exporter_excel(df_export):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Veille Législative France"

    headers = ["Date", "Média / Source", "Proposition de loi / Mesure", "Lien de consultation"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1C1917", end_color="1C1917", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D6D3CD"),
        right=Side(style="thin", color="D6D3CD"),
        top=Side(style="thin", color="D6D3CD"),
        bottom=Side(style="thin", color="D6D3CD"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row in df_export[["Date", "Source", "Titre", "Lien"]].itertuples(index=False):
        ws.append(list(row))

    for r_idx in range(2, len(df_export) + 2):
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if c_idx == 4:
                cell.font = Font(color="8B261E", underline="single")
                cell.hyperlink = cell.value

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 70)

    ws.freeze_panes = "A2"
    wb.save(output)
    return output.getvalue()


# En-tête
st.markdown(
    f"""
<div class="top-meta">
    <span>Répertoire institutionnel de veille stratégique — Presse française</span>
    <span>{datetime.now().strftime('%d %B %Y')} — Synchronisation automatique</span>
</div>
<h1 class="header-title">Présidentielle 2027<br>— <em>la course, en un coup d'œil</em></h1>
<p class="header-lead">Suivi des déclarations de candidatures dans les médias et des propositions de loi françaises (focus régulation et secteur agroalimentaire).</p>
<div class="milestones">
    <span><strong>1er tour :</strong> 18 avril 2027</span>
    <span><strong>2d tour :</strong> 2 mai 2027</span>
    <span><strong>Filtre temporel :</strong> Textes et déclarations post-1er janvier 2026</span>
</div>
<div class="main-separator"></div>
""",
    unsafe_allow_html=True,
)

tab_candidats, tab_lois = st.tabs(["Candidats déclarés", "Propositions de loi"])

# Onglet 1 : Candidats déclarés dans la presse
with tab_candidats:
    st.markdown("<h3 style='font-family:Lora,serif; font-size:1.4rem; font-weight:700;'>Personnalités déclarées candidates dans la presse</h3>", unsafe_allow_html=True)
    st.caption("Liste actualisée automatiquement d'après les dernières dépêches et déclarations publiques relevées dans les médias nationaux.")

    if data_cand:
        c1, c2 = st.columns(2)
        for idx, c in enumerate(data_cand):
            col_dest = c1 if idx % 2 == 0 else c2
            with col_dest:
                st.markdown(
                    f"""
                <div class="candidate-card">
                    <div class="candidate-name">{c['Nom']}</div>
                    <span class="candidate-party">{c['Parti']}</span>
                    <div class="candidate-quote">« {c['DerniereAnnonce']} »</div>
                    <a class="candidate-link" href="{c['Lien']}" target="_blank">Consulter l'article ({c['Source']} - {c['Date']}) ↗</a>
                </div>
                """,
                    unsafe_allow_html=True,
                )
    else:
        st.info("Aucune déclaration formelle recensée dans les médias récents.")

# Onglet 2 : Propositions de loi
with tab_lois:
    df_lois = pd.DataFrame(data_lois) if data_lois else pd.DataFrame()

    col_search, col_dl = st.columns([3, 1.2])
    with col_search:
        mot_cle = st.text_input(
            "Filtrer par mot-clé :",
            placeholder="Ex : soda, emballage, consigne, marges, alimentation, taxe...",
            label_visibility="collapsed",
        )

    if not df_lois.empty:
        if mot_cle:
            df_lois = df_lois[
                df_lois["Titre"].str.contains(mot_cle, case=False, na=False)
                | df_lois["Source"].str.contains(mot_cle, case=False, na=False)
            ]

        with col_dl:
            fichier_excel = exporter_excel(df_lois)
            st.download_button(
                label="📥 Télécharger l'Excel",
                data=fichier_excel,
                file_name=f"veille_lois_france_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

        st.caption(f"{len(df_lois)} proposition(s) de loi et mesure(s) répertoriée(s) — de la plus récente à la plus ancienne")

        st.dataframe(
            df_lois[["Date", "Source", "Titre", "Lien"]],
            column_config={
                "Lien": st.column_config.LinkColumn("Source", display_text="Consulter le texte ↗"),
                "Date": st.column_config.TextColumn("Date de parution", width="small"),
                "Source": st.column_config.TextColumn("Média français", width="small"),
                "Titre": st.column_config.TextColumn("Proposition de loi / Annonce", width="large"),
            },
            hide_index=True,
            use_container_width=True,
        )
    else:
        st.info("Aucune proposition de loi post-2026 répertoriée pour le moment.")
